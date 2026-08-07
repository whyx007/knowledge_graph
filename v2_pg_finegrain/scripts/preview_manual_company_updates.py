#!/usr/bin/env python3
"""Generate a read-only preview for manually confirmed company fields.

The source workbook has an invalid worksheet dimension (A1), so it must be
loaded in normal mode. PostgreSQL is opened with a read-only transaction. This
script never connects to Neo4j and never executes database writes.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import re
from collections import Counter
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import psycopg2
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_XLSX = ROOT / "PostgreSQL " / "企业信息确认人工确认.xlsx"
DEFAULT_OUTPUT_DIR = (
    ROOT / "data" / "update_preview" / "enterprise_info_manual_confirmation_20260807"
)
DEFAULT_DSN = "postgresql://postgres:postgres@127.0.0.1:5432/ceo_brief"


@dataclass(frozen=True)
class FieldSpec:
    excel_header: str
    db_column: str


FIELDS = (
    FieldSpec("企业简介", "company_profile"),
    FieldSpec("创始团队", "founder"),
    FieldSpec("核心技术", "core_tech"),
    FieldSpec("产品", "products"),
    FieldSpec("应用场景", "scenario"),
    FieldSpec("认证|资质|知识产权", "cert_ip"),
    FieldSpec("官网地址", "website"),
)

# This is an explicit business-approved alias, not a fuzzy match.
MANUAL_NAME_MAP = {
    "西安唐晶量子科技股份有限公司": "西安唐晶量子科技有限公司",
}

PLACEHOLDERS = {
    "-",
    "/",
    "n/a",
    "na",
    "none",
    "null",
    "不详",
    "待补充",
    "待确认",
    "无",
    "暂无",
    "未知",
    "未确认",
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="生成企业人工确认字段的只读更新预览")
    parser.add_argument("--xlsx", type=Path, default=DEFAULT_XLSX)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--dsn", default=os.getenv("PG_DSN", DEFAULT_DSN))
    parser.add_argument("--expected-matches", type=int, default=480)
    parser.add_argument("--expected-excluded", type=int, default=46)
    return parser.parse_args()


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value)
    text = text.replace("\\r\\n", " ").replace("\\n", " ").replace("\\r", " ")
    text = text.replace("\r", " ").replace("\n", " ")
    return re.sub(r"\s+", " ", text).strip()


def classify_excel_value(value: Any) -> tuple[str | None, str]:
    cleaned = clean_text(value)
    if not cleaned:
        return None, "empty"
    if cleaned.lower() in PLACEHOLDERS:
        return None, "placeholder"
    return cleaned, "valid"


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def load_excel(path: Path) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    # read_only=False is intentional: the workbook incorrectly declares A1 as
    # its dimension, while the actual rows are present in sheet2.xml.
    workbook = load_workbook(path, read_only=False, data_only=True)
    if "企业信息确认" not in workbook.sheetnames:
        raise RuntimeError("Excel 缺少工作表：企业信息确认")
    sheet = workbook["企业信息确认"]
    headers = [clean_text(sheet.cell(1, column).value) for column in range(1, sheet.max_column + 1)]
    required_headers = {"企业名称", *(field.excel_header for field in FIELDS)}
    missing_headers = sorted(required_headers - set(headers))
    if missing_headers:
        raise RuntimeError(f"Excel 缺少字段：{', '.join(missing_headers)}")

    indexes = {header: index for index, header in enumerate(headers)}
    rows: list[dict[str, Any]] = []
    blank_rows = 0
    ignored_extra_values = 0
    for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        name = clean_text(values[indexes["企业名称"]])
        if not name:
            blank_rows += 1
            continue
        record: dict[str, Any] = {"excel_row": row_number, "excel_name": name}
        for field in FIELDS:
            raw = values[indexes[field.excel_header]]
            cleaned, status = classify_excel_value(raw)
            record[field.db_column] = cleaned
            record[f"{field.db_column}_raw"] = clean_text(raw)
            record[f"{field.db_column}_status"] = status
        if "字段 1" in indexes and clean_text(values[indexes["字段 1"]]):
            ignored_extra_values += 1
        rows.append(record)

    duplicate_names = sorted(name for name, count in Counter(r["excel_name"] for r in rows).items() if count > 1)
    if duplicate_names:
        raise RuntimeError(f"Excel 企业名称重复：{duplicate_names}")

    metadata = {
        "worksheet": sheet.title,
        "physical_max_row": sheet.max_row,
        "physical_max_column": sheet.max_column,
        "company_rows": len(rows),
        "blank_name_rows": blank_rows,
        "ignored_field_1_values": ignored_extra_values,
        "headers": headers,
    }
    return rows, metadata


def load_postgres(dsn: str) -> tuple[list[dict[str, Any]], dict[str, int]]:
    columns = ["id", "name", *(field.db_column for field in FIELDS)]
    connection = psycopg2.connect(dsn)
    try:
        connection.autocommit = False
        with connection.cursor() as cursor:
            cursor.execute("SET TRANSACTION READ ONLY")
            cursor.execute(
                "SELECT " + ", ".join(columns) + " FROM public.companies ORDER BY id"
            )
            records = [dict(zip(columns, row)) for row in cursor.fetchall()]
            cursor.execute(
                "SELECT "
                "count(*) FILTER (WHERE core_tech_vec IS NOT NULL), "
                "count(*) FILTER (WHERE full_text_vec IS NOT NULL) "
                "FROM public.companies"
            )
            core_tech_vectors, full_text_vectors = cursor.fetchone()
        connection.rollback()
        return records, {
            "core_tech_vectors": core_tech_vectors,
            "full_text_vectors": full_text_vectors,
        }
    finally:
        connection.close()


def website_review(value: str) -> tuple[str, str]:
    lowered = value.lower()
    if lowered.startswith(("https://", "http://")):
        return "ready", value
    if lowered.startswith("www."):
        return "www_without_protocol", f"https://{value}"
    if re.fullmatch(r"(?:[a-z0-9-]+\.)+[a-z]{2,}(?:/[^\s]*)?", lowered):
        return "bare_domain_without_protocol", f"https://{value}"
    return "needs_manual_review", value


def build_preview(
    excel_rows: list[dict[str, Any]],
    pg_rows: list[dict[str, Any]],
) -> dict[str, Any]:
    pg_by_name = {row["name"]: row for row in pg_rows}
    if len(pg_by_name) != len(pg_rows):
        raise RuntimeError("PostgreSQL companies.name 不是唯一值")

    companies: list[dict[str, Any]] = []
    field_details: list[dict[str, Any]] = []
    staging_rows: list[dict[str, Any]] = []
    excluded: list[dict[str, Any]] = []
    websites: list[dict[str, Any]] = []
    matched_pg_ids: set[int] = set()

    for excel in excel_rows:
        excel_name = excel["excel_name"]
        db_name = MANUAL_NAME_MAP.get(excel_name, excel_name)
        match_method = "manual_mapping" if excel_name in MANUAL_NAME_MAP else "exact"
        pg = pg_by_name.get(db_name)
        if pg is None:
            excluded.append(
                {
                    "excel_row": excel["excel_row"],
                    "excel_name": excel_name,
                    "reason": "not_in_existing_postgresql",
                }
            )
            continue
        if pg["id"] in matched_pg_ids:
            raise RuntimeError(f"多个 Excel 企业映射到同一 PostgreSQL id：{pg['id']}")
        matched_pg_ids.add(pg["id"])

        changed_fields: list[str] = []
        review_fields: list[str] = []
        unchanged_fields: list[str] = []
        preserved_fields: list[str] = []
        staging: dict[str, Any] = {
            "company_id": pg["id"],
            "postgres_name": pg["name"],
            "excel_name": excel_name,
            "match_method": match_method,
        }
        for field in FIELDS:
            column = field.db_column
            source_status = excel[f"{column}_status"]
            proposed = excel[column]
            old_raw = pg[column]
            old_clean = clean_text(old_raw)
            if source_status != "valid":
                action = "preserve"
                reason = "excel_empty" if source_status == "empty" else "excel_placeholder"
                effective = old_raw
                preserved_fields.append(column)
            elif proposed == old_clean:
                action = "unchanged"
                reason = "same_after_whitespace_normalization"
                effective = old_raw
                unchanged_fields.append(column)
            elif column == "website" and website_review(proposed)[0] != "ready":
                action = "review"
                reason = website_review(proposed)[0]
                effective = old_raw
                review_fields.append(column)
            else:
                action = "update"
                reason = "valid_confirmed_value_differs"
                effective = proposed
                changed_fields.append(column)

            staging[column] = effective
            staging[f"update_{column}"] = action == "update"
            field_details.append(
                {
                    "company_id": pg["id"],
                    "postgres_name": pg["name"],
                    "excel_name": excel_name,
                    "match_method": match_method,
                    "db_field": column,
                    "excel_header": field.excel_header,
                    "action": action,
                    "reason": reason,
                    "original_value": old_raw,
                    "excel_value": excel[f"{column}_raw"],
                    "proposed_value": effective,
                }
            )

        companies.append(
            {
                "company_id": pg["id"],
                "postgres_name": pg["name"],
                "excel_name": excel_name,
                "match_method": match_method,
                "changed_field_count": len(changed_fields),
                "review_field_count": len(review_fields),
                "changed_fields": ",".join(changed_fields),
                "review_fields": ",".join(review_fields),
                "unchanged_fields": ",".join(unchanged_fields),
                "preserved_fields": ",".join(preserved_fields),
            }
        )
        staging_rows.append(staging)

        website = excel["website"]
        if website:
            status, suggested = website_review(website)
            if status != "ready":
                websites.append(
                    {
                        "company_id": pg["id"],
                        "postgres_name": pg["name"],
                        "excel_website": website,
                        "review_status": status,
                        "suggested_value": suggested,
                        "note": "仅为预览建议，验证前不应用",
                    }
                )

    missing_pg = [
        {"company_id": row["id"], "postgres_name": row["name"], "reason": "missing_from_excel"}
        for row in pg_rows
        if row["id"] not in matched_pg_ids
    ]
    return {
        "companies": sorted(companies, key=lambda row: row["company_id"]),
        "field_details": sorted(
            field_details, key=lambda row: (row["company_id"], row["db_field"])
        ),
        "staging_rows": sorted(staging_rows, key=lambda row: row["company_id"]),
        "excluded": sorted(excluded, key=lambda row: row["excel_row"]),
        "websites": sorted(websites, key=lambda row: row["company_id"]),
        "missing_pg": missing_pg,
    }


def assert_expected(preview: dict[str, Any], expected_matches: int, expected_excluded: int) -> None:
    actual_matches = len(preview["companies"])
    actual_excluded = len(preview["excluded"])
    if actual_matches != expected_matches:
        raise RuntimeError(f"匹配企业数断言失败：expected={expected_matches}, actual={actual_matches}")
    if actual_excluded != expected_excluded:
        raise RuntimeError(f"排除企业数断言失败：expected={expected_excluded}, actual={actual_excluded}")
    if preview["missing_pg"]:
        raise RuntimeError(f"有 {len(preview['missing_pg'])} 家 PostgreSQL 企业未被 Excel 覆盖")


def write_csv(path: Path, rows: list[dict[str, Any]], fieldnames: list[str]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as output:
        writer = csv.DictWriter(output, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def add_sheet(workbook: Workbook, title: str, rows: list[dict[str, Any]], columns: list[str]) -> None:
    sheet = workbook.create_sheet(title)
    sheet.append(columns)
    for row in rows:
        values = []
        for column in columns:
            value = row.get(column)
            if isinstance(value, (dict, list, tuple, set)):
                value = json.dumps(value, ensure_ascii=False, sort_keys=True)
            values.append(value)
        sheet.append(values)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column_cells in sheet.columns:
        max_length = max(len(str(cell.value or "")) for cell in column_cells[:200])
        sheet.column_dimensions[column_cells[0].column_letter].width = min(max(max_length + 2, 12), 50)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def write_outputs(
    output_dir: Path,
    preview: dict[str, Any],
    summary: dict[str, Any],
) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    company_columns = [
        "company_id", "postgres_name", "excel_name", "match_method",
        "changed_field_count", "review_field_count", "changed_fields", "review_fields",
        "unchanged_fields", "preserved_fields",
    ]
    detail_columns = [
        "company_id", "postgres_name", "excel_name", "match_method", "db_field",
        "excel_header", "action", "reason", "original_value", "excel_value", "proposed_value",
    ]
    excluded_columns = ["excel_row", "excel_name", "reason"]
    website_columns = [
        "company_id", "postgres_name", "excel_website", "review_status",
        "suggested_value", "note",
    ]
    staging_columns = ["company_id", "postgres_name", "excel_name", "match_method"]
    for field in FIELDS:
        staging_columns.extend([field.db_column, f"update_{field.db_column}"])

    write_csv(output_dir / "company_update_preview.csv", preview["companies"], company_columns)
    write_csv(output_dir / "field_level_diff.csv", preview["field_details"], detail_columns)
    write_csv(output_dir / "excluded_companies.csv", preview["excluded"], excluded_columns)
    write_csv(output_dir / "website_review.csv", preview["websites"], website_columns)
    write_csv(output_dir / "update_staging.csv", preview["staging_rows"], staging_columns)
    (output_dir / "summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )

    workbook = Workbook()
    workbook.remove(workbook.active)
    add_sheet(
        workbook,
        "概要",
        [{"指标": key, "值": value} for key, value in summary.items()],
        ["指标", "值"],
    )
    add_sheet(workbook, "企业更新预览", preview["companies"], company_columns)
    add_sheet(workbook, "字段级差异", preview["field_details"], detail_columns)
    add_sheet(workbook, "排除清单", preview["excluded"], excluded_columns)
    add_sheet(workbook, "官网待复核", preview["websites"], website_columns)
    workbook.save(output_dir / "manual_company_update_preview.xlsx")


def main() -> None:
    args = parse_args()
    excel_path = args.xlsx.resolve()
    excel_rows, excel_metadata = load_excel(excel_path)
    pg_rows, pg_metadata = load_postgres(args.dsn)
    preview = build_preview(excel_rows, pg_rows)
    assert_expected(preview, args.expected_matches, args.expected_excluded)

    action_counts = Counter(row["action"] for row in preview["field_details"])
    changed_by_field = Counter(
        row["db_field"] for row in preview["field_details"] if row["action"] == "update"
    )
    differences_by_field = Counter(
        row["db_field"]
        for row in preview["field_details"]
        if row["action"] in {"update", "review"}
    )
    website_review_counts = Counter(row["review_status"] for row in preview["websites"])
    generated_at = datetime.now(timezone.utc).isoformat(timespec="seconds")
    summary = {
        "generated_at_utc": generated_at,
        "source_file": str(excel_path),
        "source_sha256": sha256_file(excel_path),
        "database": "ceo_brief/public.companies (read-only transaction)",
        "excel_company_rows": len(excel_rows),
        "excel_blank_name_rows": excel_metadata["blank_name_rows"],
        "ignored_field_1_values": excel_metadata["ignored_field_1_values"],
        "postgres_companies": len(pg_rows),
        "matched_companies": len(preview["companies"]),
        "exact_matches": sum(row["match_method"] == "exact" for row in preview["companies"]),
        "manual_matches": sum(row["match_method"] == "manual_mapping" for row in preview["companies"]),
        "excluded_companies": len(preview["excluded"]),
        "postgres_companies_missing_from_excel": len(preview["missing_pg"]),
        "companies_with_differences": sum(
            row["changed_field_count"] + row["review_field_count"] > 0
            for row in preview["companies"]
        ),
        "companies_with_ready_updates": sum(
            row["changed_field_count"] > 0 for row in preview["companies"]
        ),
        "companies_with_core_tech_or_products_changes": len(
            {
                row["company_id"]
                for row in preview["field_details"]
                if row["action"] == "update" and row["db_field"] in {"core_tech", "products"}
            }
        ),
        "field_updates": action_counts["update"],
        "field_reviews_pending": action_counts["review"],
        "field_unchanged": action_counts["unchanged"],
        "field_preserved": action_counts["preserve"],
        "website_values_needing_review": len(preview["websites"]),
        "website_www_without_protocol": website_review_counts["www_without_protocol"],
        "website_bare_domain_without_protocol": website_review_counts[
            "bare_domain_without_protocol"
        ],
        "website_values_needing_manual_review": website_review_counts["needs_manual_review"],
        "existing_core_tech_vectors": pg_metadata["core_tech_vectors"],
        "existing_full_text_vectors": pg_metadata["full_text_vectors"],
        "updates_by_field": dict(sorted(changed_by_field.items())),
        "differences_by_field": dict(sorted(differences_by_field.items())),
        "database_writes_performed": False,
        "neo4j_connections_performed": False,
    }
    write_outputs(args.output_dir.resolve(), preview, summary)
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    print(f"\n预览已生成：{args.output_dir.resolve()}")


if __name__ == "__main__":
    main()
