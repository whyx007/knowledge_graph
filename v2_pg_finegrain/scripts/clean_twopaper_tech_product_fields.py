#!/usr/bin/env python3
"""Clean twopaper_v2.xlsx core-tech and product/application columns."""

from __future__ import annotations

import re
import shutil
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
XLSX_PATH = ROOT / "PostgreSQL " / "twopaper_v2.xlsx"
REPORT_PATH = ROOT / "docs" / "twopaper_v2_核心技术与产品字段清洗报告.md"

CORE_COL = "核心技术"
PRODUCT_COL = "产品及应用领域"

UNRELATED_RE = re.compile(
    r"(融资|投资|估值|估价|营收|收入|回款|订单|客户|一供|独供|供应商地位|商业化进展|商业进展|"
    r"投后|股权|持股|上市|IPO|获投|领投|跟投|资金支持|政府资源|市场验证|规模化生产经验|"
    r"员工规模|人员规模|短板|风险|挑战|竞争激烈|头部客户|批量导入|已导入|标杆客户|合作客户|"
    r"市场情况|市场装机量|已获得投资情况|Pre-?A|A\+*\s*轮|B\+*\s*轮|C\+*\s*轮|天使轮|种子轮|"
    r"科创板|中小企业发展基金|基石资本|东方富海|资本|本土团队|创始团队|核心团队|核心专家)"
)

CORE_SIGNAL_RE = re.compile(
    r"(核心技术|技术|工艺|算法|模型|架构|芯片|材料|传感|激光|光学|控制|软件|平台|系统|"
    r"制备|制造|检测|测量|封装|设计|研发|开发|自主|专利|壁垒|能力|调控|合成|聚合|成像|通信|"
    r"计算|训练|推理|编译|量产工艺|解决方案技术)"
)

PRODUCT_SIGNAL_RE = re.compile(
    r"(产品|应用|应用场景|领域|设备|装备|仪器|模组|模块|器件|零部件|材料|膜片|元件|整机|"
    r"服务|方案|产线|电池|激光器|光源|泵|相机|雷达|机器人|平台|软件|系统|终端|用于|面向|"
    r"覆盖|适配|包括|涵盖|提供)"
)

HEADING_ONLY_RE = re.compile(
    r"^(核心技术|产品|产品简介|产品及解决方案简介|产品及应用领域|应用场景|主要产品|公司产品|核心产品|产品线|主要业务)[:：]?$"
)


def clean_text(value: object) -> str:
    if value is None:
        return ""
    text = str(value).replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t\u3000]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def normalize_piece(value: str) -> str:
    return re.sub(r"[\s，,。；;：:、（）()·.]+", "", value).lower()


def split_pieces(text: str) -> list[str]:
    text = clean_text(text)
    if not text:
        return []
    raw_lines: list[str] = []
    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        # Split only at strong section boundaries. Keep comma/semicolon detail inside a sentence.
        line = re.sub(r"(?<![A-Za-z0-9])([0-9]+[、)])", r"\n\1", line)
        line = re.sub(r"(?<!公司)(?<!主要)(?<!核心)(产品[0-9一二三四五六七八九十]+[:：])", r"\n\1", line)
        line = re.sub(r"(应用场景[:：]|核心技术[:：]|商业化进展[:：]?)", r"\n\1", line)
        raw_lines.extend(part.strip() for part in line.splitlines() if part.strip())
    return raw_lines


def strip_prefix(piece: str) -> str:
    text = piece.strip()
    text = re.sub(r"^[0-9]+[、)]\s*", "", text)
    text = re.sub(r"^[0-9]+[.．](?!\d)\s*", "", text)
    text = re.sub(r"^[0-9]+[.．](?=[A-Za-z\u4e00-\u9fff])", "", text)
    text = re.sub(r"^[0-9]+[.．](?=[0-9]+[A-Za-z])", "", text)
    text = re.sub(r"^(核心技术|技术核心|产品及应用领域|产品及解决方案简介|产品线|核心产品|主要产品|公司产品|应用领域|应用场景|产品简介|产品)[:：]\s*", "", text)
    text = re.sub(r"^.*主要产品[:：]?\s*$", "", text)
    if "|" in text:
        parts = [part.strip() for part in text.split("|") if part.strip()]
        parts = [
            part
            for part in parts
            if not re.search(r"(Pre-?A|[ABC]\+*\s*轮|天使轮|种子轮|北京|上海|深圳|西安|苏州|杭州)$", part)
        ]
        text = " | ".join(parts)
    return text.strip(" 。；;")


def is_unrelated(piece: str) -> bool:
    text = strip_prefix(piece)
    if not text or HEADING_ONLY_RE.match(text):
        return True
    if re.fullmatch(r"[0-9一二三四五六七八九十]+", text):
        return True
    if UNRELATED_RE.search(text):
        # Keep explicit technical cost/performance descriptions.
        if re.search(r"(成本下降|成本降低|低成本).*(技术|工艺|方案|设计)", text):
            return False
        return True
    if re.search(r"^(公司简介|企业简介)[:：]", text):
        return True
    if re.search(r"(国内唯一|行业第一|第一家|首家|头部|领先地位|市场占有|量产出货|出货量|已完成.*导入)", text):
        if not re.search(r"(产品|设备|系统|应用|用于|面向|技术|工艺|材料|芯片|器件)", text):
            return True
    return False


def classify(piece: str, source: str) -> str:
    text = strip_prefix(piece)
    if is_unrelated(text):
        return "drop"
    core = bool(CORE_SIGNAL_RE.search(text))
    product = bool(PRODUCT_SIGNAL_RE.search(text))
    has_product_prefix = bool(re.match(r"^(产品|应用场景|主要产品|公司产品|核心产品|产品线)[:：]", piece.strip()))
    has_core_prefix = bool(re.match(r"^(核心技术|技术核心)[:：]", piece.strip()))

    if has_core_prefix:
        return "core"
    if has_product_prefix:
        return "product"
    if source == "core":
        if product and not core:
            return "product"
        return "core"
    if core and not product:
        return "core"
    return "product"


def append_unique(target: list[str], piece: str) -> bool:
    text = strip_prefix(piece)
    if not text:
        return False
    norm = normalize_piece(text)
    if not norm:
        return False
    for existing in target:
        existing_norm = normalize_piece(existing)
        if norm == existing_norm:
            return False
        if len(norm) > 18 and norm in existing_norm:
            return False
        if len(existing_norm) > 18 and existing_norm in norm:
            target[target.index(existing)] = text
            return True
    target.append(text)
    return True


def join_pieces(pieces: list[str]) -> str:
    cleaned = []
    for piece in pieces:
        text = strip_prefix(piece)
        text = re.sub(r"\s+", " ", text).strip(" 。；;")
        if text:
            cleaned.append(text)
    if not cleaned:
        return ""
    if len(cleaned) == 1:
        return cleaned[0]
    return "\n".join(f"{idx}. {piece}" for idx, piece in enumerate(cleaned, 1))


def md_cell(value: str, limit: int = 180) -> str:
    text = re.sub(r"\s+", " ", value or "").strip()
    if len(text) > limit:
        text = text[: limit - 1] + "…"
    return text.replace("|", "\\|")


def main() -> None:
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_path = XLSX_PATH.with_name(f"twopaper_v2_backup_before_field_clean_{timestamp}.xlsx")
    shutil.copy2(XLSX_PATH, backup_path)

    wb = load_workbook(XLSX_PATH)
    ws = wb["企业信息提取"]
    headers = [cell.value for cell in ws[1]]
    core_idx = headers.index(CORE_COL) + 1
    product_idx = headers.index(PRODUCT_COL) + 1
    name_idx = headers.index("公司名称") + 1

    changes: list[dict[str, str]] = []
    stats = defaultdict(int)

    for row_idx in range(2, ws.max_row + 1):
        name = clean_text(ws.cell(row_idx, name_idx).value)
        if not name:
            continue
        old_core = clean_text(ws.cell(row_idx, core_idx).value)
        old_product = clean_text(ws.cell(row_idx, product_idx).value)
        core_parts: list[str] = []
        product_parts: list[str] = []
        dropped: list[str] = []

        for source, text in (("core", old_core), ("product", old_product)):
            for piece in split_pieces(text):
                category = classify(piece, source)
                if category == "drop":
                    stripped = strip_prefix(piece)
                    if stripped:
                        dropped.append(stripped)
                    continue
                if category == "core":
                    append_unique(core_parts, piece)
                else:
                    append_unique(product_parts, piece)

        new_core = join_pieces(core_parts)
        new_product = join_pieces(product_parts)

        if new_core != old_core or new_product != old_product:
            ws.cell(row_idx, core_idx).value = new_core or None
            ws.cell(row_idx, product_idx).value = new_product or None
            stats["changed_rows"] += 1
            if old_core != new_core:
                stats["core_changed"] += 1
            if old_product != new_product:
                stats["product_changed"] += 1
            if dropped:
                stats["rows_with_drops"] += 1
            changes.append(
                {
                    "name": name,
                    "old_core": old_core,
                    "new_core": new_core,
                    "old_product": old_product,
                    "new_product": new_product,
                    "dropped": "；".join(dropped),
                }
            )

    wb.save(XLSX_PATH)

    lines = [
        "# twopaper_v2.xlsx 核心技术与产品字段清洗报告",
        "",
        f"- 生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
        f"- 清洗文件：`{XLSX_PATH}`",
        f"- 备份文件：`{backup_path}`",
        f"- 变更行数：{stats['changed_rows']}",
        f"- `核心技术` 变更行数：{stats['core_changed']}",
        f"- `产品及应用领域` 变更行数：{stats['product_changed']}",
        f"- 删除无关内容行数：{stats['rows_with_drops']}",
        "",
        "## 清洗口径",
        "",
        "- `核心技术`：保留技术、工艺、算法、模型、材料、芯片、系统架构、研发能力等内容。",
        "- `产品及应用领域`：保留产品、设备、模组、材料、服务、解决方案、应用场景和面向领域等内容。",
        "- 删除内容：投资融资、估值营收、客户导入、订单、一供/独供、商业化进展、团队短板/风险等与两列标题无关的信息。",
        "",
        "## 变更明细",
        "",
        "| 公司名称 | 原核心技术 | 新核心技术 | 原产品及应用领域 | 新产品及应用领域 | 删除/剔除内容 |",
        "|---|---|---|---|---|---|",
    ]
    for change in changes:
        lines.append(
            "| "
            + " | ".join(
                [
                    md_cell(change["name"], 80),
                    md_cell(change["old_core"]),
                    md_cell(change["new_core"]),
                    md_cell(change["old_product"]),
                    md_cell(change["new_product"]),
                    md_cell(change["dropped"]),
                ]
            )
            + " |"
        )
    REPORT_PATH.write_text("\n".join(lines) + "\n", encoding="utf-8")

    print(f"backup={backup_path}")
    print(f"report={REPORT_PATH}")
    print(f"changed_rows={stats['changed_rows']}")
    print(f"core_changed={stats['core_changed']}")
    print(f"product_changed={stats['product_changed']}")
    print(f"rows_with_drops={stats['rows_with_drops']}")


if __name__ == "__main__":
    main()
