#!/usr/bin/env python3
"""Merge twopaper_v2.xlsx company intro and founder fields into PostgreSQL."""

from __future__ import annotations

import difflib
import re
from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

import psycopg2
from openpyxl import load_workbook
from psycopg2 import sql
from psycopg2.extras import RealDictCursor


ROOT = Path(__file__).resolve().parents[1]
XLSX_PATH = ROOT / "PostgreSQL " / "twopaper_v2.xlsx"
MATCH_REPORT = ROOT / "docs" / "twopaper_v2_companies匹配报告.md"
FOUNDER_REPORT = ROOT / "docs" / "twopaper_v2_founder融合更新报告.md"
DSN = "postgresql://postgres:postgres@127.0.0.1:5432/ceo_brief"
INTRO_COLUMN = "company_profile"


LEGAL_SUFFIXES = (
    "股份有限公司",
    "有限责任公司",
    "集团有限公司",
    "控股有限公司",
    "科技股份有限公司",
    "科技有限公司",
    "技术股份有限公司",
    "技术有限公司",
    "有限公司",
)


@dataclass(frozen=True)
class Match:
    excel_name: str
    db_id: int
    db_name: str
    method: str
    confidence: str


def clean_text(value: object) -> str:
    if value is None:
        return ""
    text = str(value).replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t\u3000]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def one_line(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip()


def md_cell(value: object, limit: int | None = None) -> str:
    text = "" if value is None else one_line(str(value))
    if limit and len(text) > limit:
        text = text[: limit - 1] + "…"
    return text.replace("|", "\\|").replace("\n", "<br>")


def normalize_name(value: str) -> str:
    text = clean_text(value)
    text = text.replace("（", "(").replace("）", ")")
    text = re.sub(r"[-－—–](已退出|退出)$", "", text)
    text = re.sub(r"(已退出|退出)$", "", text)
    text = re.sub(r"[\s\u3000]+", "", text)
    return text


def legal_core(value: str) -> str:
    text = normalize_name(value)
    for suffix in LEGAL_SUFFIXES:
        if text.endswith(suffix) and len(text) > len(suffix) + 2:
            return text[: -len(suffix)]
    return text


def split_company_names(value: str) -> list[str]:
    parts = re.split(r"[、,，/]", clean_text(value))
    return [part.strip() for part in parts if part and part.strip()]


def founder_body(value: str) -> str:
    text = one_line(clean_text(value))
    text = re.sub(r"^(补充信息|创始人/创始团队|创始团队|核心团队|团队优势|团队包括|核心创始团队)[:：]\s*", "", text)
    text = re.sub(r"\s+", " ", text)
    return dedupe_founder_text(text).strip(" 。；;")


def trim_non_founder_sections(value: str) -> str:
    text = one_line(value)
    cut_markers = (
        "主要产品：",
        "主要产品:",
        "主要业务：",
        "主要业务:",
        "产品线",
        "应用场景：",
        "应用场景:",
        "商业化进展",
    )
    founder_markers = ("核心团队", "创始团队", "创始人", "联合创始人", "董事长", "首席科学家")
    for marker in cut_markers:
        pos = text.find(marker)
        if pos < 25:
            continue
        tail = text[pos + len(marker) :]
        next_positions = [tail.find(founder_marker) for founder_marker in founder_markers if tail.find(founder_marker) >= 0]
        if next_positions:
            keep_from = pos + len(marker) + min(next_positions)
            text = f"{text[:pos].strip(' 。；;')} {text[keep_from:].strip()}"
        else:
            text = text[:pos].strip(" 。；;")
    return text


def dedupe_founder_text(value: str) -> str:
    text = trim_non_founder_sections(value)
    text = text.replace("；", "。").replace(";", "。")
    chunks = [chunk.strip(" 。") for chunk in re.split(r"(?<=[。!?！？])\s*|\n+", text) if chunk.strip(" 。")]
    if not chunks:
        chunks = [chunk.strip(" 。") for chunk in re.split(r"\s{2,}", text) if chunk.strip(" 。")]

    kept: list[str] = []
    kept_norms: list[str] = []
    for chunk in chunks:
        norm = normalized_text(chunk)
        if not norm:
            continue
        duplicate = False
        for existing_norm in kept_norms:
            if norm == existing_norm:
                duplicate = True
                break
            if len(norm) > 18 and norm in existing_norm:
                duplicate = True
                break
            if len(existing_norm) > 18 and existing_norm in norm:
                idx = kept_norms.index(existing_norm)
                kept[idx] = chunk
                kept_norms[idx] = norm
                duplicate = True
                break
        if not duplicate:
            kept.append(chunk)
            kept_norms.append(norm)

    merged = "。".join(kept)
    merged = re.sub(r"([0-9])\.。", r"\1. ", merged)
    merged = re.sub(r"(核心团队\s*){2,}", "核心团队 ", merged)
    merged = re.sub(r"(创始团队\s*){2,}", "创始团队 ", merged)
    merged = re.sub(r"核心团队\s+核心团队", "核心团队", merged)
    merged = re.sub(r"创始团队\s+创始团队", "创始团队", merged)
    return merged.strip(" 。；;")


def normalized_text(value: str) -> str:
    return re.sub(r"[\s，,。；;：:、（）()·.]+", "", clean_text(value)).lower()


def merge_founder(existing: str, addition: str) -> tuple[str, bool]:
    old_body = founder_body(existing)
    new_body = founder_body(addition)
    if not new_body:
        return clean_text(existing), False
    if not old_body:
        return f"创始人/创始团队：{new_body}", True

    old_norm = normalized_text(old_body)
    new_norm = normalized_text(new_body)
    similarity = difflib.SequenceMatcher(None, old_norm, new_norm).ratio() if old_norm and new_norm else 0.0
    overlap = 0.0
    if old_norm and new_norm:
        common = 0
        for token in re.findall(r"[\u4e00-\u9fffA-Za-z0-9]{2,}", old_body):
            token_norm = normalized_text(token)
            if len(token_norm) >= 2 and token_norm in new_norm:
                common += len(token_norm)
        overlap = common / max(len(old_norm), 1)

    if not new_norm or new_norm in old_norm:
        merged = f"创始人/创始团队：{old_body}"
    elif old_norm and old_norm in new_norm:
        merged = f"创始人/创始团队：{new_body}"
    elif similarity >= 0.42 or overlap >= 0.35:
        merged = f"创始人/创始团队：{new_body if len(new_norm) >= len(old_norm) else old_body}"
    else:
        merged = f"创始人/创始团队：{old_body}。{new_body}"

    merged = "创始人/创始团队：" + dedupe_founder_text(re.sub(r"^创始人/创始团队[:：]\s*", "", merged))
    merged = re.sub(r"。{2,}", "。", merged).strip()
    return merged, normalized_text(merged) != normalized_text(existing)


def load_excel_rows() -> list[dict[str, str]]:
    wb = load_workbook(XLSX_PATH, read_only=True, data_only=True)
    ws = wb["企业信息提取"]
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    rows: list[dict[str, str]] = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        row = {key: clean_text(value) for key, value in zip(headers, values)}
        if row.get("公司名称"):
            rows.append(row)
    return rows


def unique_core_index(companies: list[dict]) -> dict[str, dict]:
    buckets: dict[str, list[dict]] = defaultdict(list)
    for company in companies:
        buckets[legal_core(company["name"])].append(company)
    return {key: values[0] for key, values in buckets.items() if len(values) == 1}


def find_matches(rows: list[dict[str, str]], companies: list[dict]) -> tuple[list[tuple[dict, Match]], list[dict]]:
    by_name = {company["name"]: company for company in companies}
    by_norm = {normalize_name(company["name"]): company for company in companies}
    by_core = unique_core_index(companies)
    matched: list[tuple[dict, Match]] = []
    unmatched: list[dict] = []
    seen_pairs: set[tuple[str, int]] = set()

    for row in rows:
        excel_name = row["公司名称"]
        candidates: list[Match] = []
        if excel_name in by_name:
            c = by_name[excel_name]
            candidates.append(Match(excel_name, c["id"], c["name"], "exact", "high"))
        elif normalize_name(excel_name) in by_norm:
            c = by_norm[normalize_name(excel_name)]
            candidates.append(Match(excel_name, c["id"], c["name"], "normalized_parentheses_space", "high"))
        else:
            for part in split_company_names(excel_name):
                if normalize_name(part) in by_norm:
                    c = by_norm[normalize_name(part)]
                    candidates.append(Match(part, c["id"], c["name"], "split_exact", "high"))
            if not candidates and legal_core(excel_name) in by_core:
                c = by_core[legal_core(excel_name)]
                candidates.append(Match(excel_name, c["id"], c["name"], "legal_suffix", "high"))

        if candidates:
            for candidate in candidates:
                pair = (excel_name, candidate.db_id)
                if pair not in seen_pairs:
                    matched.append((row, candidate))
                    seen_pairs.add(pair)
        else:
            unmatched.append(row)

    return matched, unmatched


def top_fuzzy(row: dict[str, str], companies: list[dict]) -> tuple[int, str, float] | None:
    source = normalize_name(row["公司名称"])
    best: tuple[float, int, str] | None = None
    for company in companies:
        target = normalize_name(company["name"])
        score = difflib.SequenceMatcher(None, source, target).ratio()
        if best is None or score > best[0]:
            best = (score, company["id"], company["name"])
    if best and best[0] >= 0.88:
        return best[1], best[2], best[0]
    return None


def ensure_intro_column(cur) -> None:
    cur.execute(
        """
        SELECT 1
        FROM information_schema.columns
        WHERE table_schema = 'public'
          AND table_name = 'companies'
          AND column_name = %s
        """,
        (INTRO_COLUMN,),
    )
    if cur.fetchone():
        return
    cur.execute(sql.SQL("ALTER TABLE companies ADD COLUMN {} text").format(sql.Identifier(INTRO_COLUMN)))
    cur.execute(
        sql.SQL("COMMENT ON COLUMN companies.{} IS %s").format(sql.Identifier(INTRO_COLUMN)),
        ("twopaper_v2.xlsx 匹配导入的企业简介",),
    )


def create_backup(cur, backup_table: str) -> None:
    cur.execute(
        sql.SQL(
            """
            CREATE TABLE {} AS
            SELECT id, name, founder, {} AS company_intro_before
            FROM companies
            """
        ).format(sql.Identifier(backup_table), sql.Identifier(INTRO_COLUMN))
    )


def write_match_report(
    rows: list[dict[str, str]],
    matched: list[tuple[dict, Match]],
    unmatched: list[dict],
    fuzzy: dict[str, tuple[int, str, float]],
    backup_table: str,
) -> None:
    method_counts: dict[str, int] = defaultdict(int)
    for _, match in matched:
        method_counts[match.method] += 1

    lines = [
        "# twopaper_v2.xlsx 与 PostgreSQL companies 匹配报告",
        "",
        f"- 生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
        f"- Excel 文件：`{XLSX_PATH}`",
        "- 数据库：`ceo_brief.public.companies`",
        f"- 更新前备份表：`{backup_table}`",
        f"- Excel 有效公司行数：{len(rows)}",
        f"- 高置信匹配记录数：{len(matched)}",
        f"- 未写库 Excel 行数：{len(unmatched)}",
        "",
        "## 匹配方式统计",
        "",
        "| 匹配方式 | 数量 |",
        "|---|---:|",
    ]
    for method in sorted(method_counts):
        lines.append(f"| {method} | {method_counts[method]} |")

    lines.extend(["", "## 已写库高置信匹配", "", "| Excel 公司名称 | DB id | DB 公司名称 | 匹配方式 | 企业简介 | 创始人/团队 |", "|---|---:|---|---|---|---|"])
    for row, match in sorted(matched, key=lambda item: item[1].db_id):
        lines.append(
            "| "
            + " | ".join(
                [
                    md_cell(row["公司名称"], 80),
                    str(match.db_id),
                    md_cell(match.db_name, 80),
                    match.method,
                    "有" if row.get("公司简介") else "无",
                    "有" if row.get("创始人/创始团队") else "无",
                ]
            )
            + " |"
        )

    lines.extend(["", "## 未写库记录", "", "| Excel 公司名称 | 可能相似 DB 记录 | 说明 |", "|---|---|---|"])
    for row in unmatched:
        possible = fuzzy.get(row["公司名称"])
        if possible:
            candidate = f"ENT_PG_{possible[0]} / {possible[1]} / score={possible[2]:.2f}"
            note = "仅作为人工复核候选，未自动更新"
        else:
            candidate = ""
            note = "未找到高置信匹配"
        lines.append(f"| {md_cell(row['公司名称'], 100)} | {md_cell(candidate, 120)} | {note} |")

    MATCH_REPORT.write_text("\n".join(lines) + "\n", encoding="utf-8")


def write_founder_report(changes: list[dict], skipped: list[tuple[dict, Match]]) -> None:
    lines = [
        "# twopaper_v2.xlsx 创始人字段融合更新报告",
        "",
        f"- 生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
        "- 更新字段：`companies.founder`",
        f"- 实际发生 founder 更新：{len(changes)} 条",
        f"- 高置信匹配但 Excel 创始人/创始团队为空：{len(skipped)} 条",
        "",
        "## 更新明细",
        "",
        "| DB id | 公司名称 | 匹配方式 | 原 founder 摘要 | Excel 创始人/创始团队 | 更新后 founder 摘要 |",
        "|---:|---|---|---|---|---|",
    ]
    for change in sorted(changes, key=lambda item: item["db_id"]):
        lines.append(
            "| "
            + " | ".join(
                [
                    str(change["db_id"]),
                    md_cell(change["db_name"], 80),
                    change["method"],
                    md_cell(change["old_founder"], 180),
                    md_cell(change["excel_founder"], 220),
                    md_cell(change["new_founder"], 260),
                ]
            )
            + " |"
        )

    lines.extend(["", "## 高置信匹配但未更新 founder", "", "| DB id | 公司名称 | 匹配方式 | 原因 |", "|---:|---|---|---|"])
    for _, match in sorted(skipped, key=lambda item: item[1].db_id):
        lines.append(f"| {match.db_id} | {md_cell(match.db_name, 80)} | {match.method} | Excel 创始人/创始团队为空 |")

    FOUNDER_REPORT.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> None:
    rows = load_excel_rows()
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_table = f"companies_twopaper_backup_{timestamp}"

    conn = psycopg2.connect(DSN)
    try:
        with conn:
            with conn.cursor(cursor_factory=RealDictCursor) as cur:
                ensure_intro_column(cur)
                create_backup(cur, backup_table)

                cur.execute("SELECT id, name, founder FROM companies ORDER BY id")
                companies = list(cur.fetchall())
                matched, unmatched = find_matches(rows, companies)
                fuzzy = {row["公司名称"]: top_fuzzy(row, companies) for row in unmatched}
                fuzzy = {key: value for key, value in fuzzy.items() if value is not None}

                current_by_id = {company["id"]: company for company in companies}
                founder_changes: list[dict] = []
                founder_skipped: list[tuple[dict, Match]] = []

                for row, match in matched:
                    intro = clean_text(row.get("公司简介", ""))
                    excel_founder = clean_text(row.get("创始人/创始团队", ""))
                    current = current_by_id[match.db_id]
                    old_founder = clean_text(current.get("founder", ""))
                    new_founder, founder_changed = merge_founder(old_founder, excel_founder)

                    if excel_founder and founder_changed:
                        founder_changes.append(
                            {
                                "db_id": match.db_id,
                                "db_name": match.db_name,
                                "method": match.method,
                                "old_founder": old_founder,
                                "excel_founder": excel_founder,
                                "new_founder": new_founder,
                            }
                        )
                    elif not excel_founder:
                        founder_skipped.append((row, match))

                    cur.execute(
                        sql.SQL("UPDATE companies SET {} = %s, founder = %s WHERE id = %s").format(
                            sql.Identifier(INTRO_COLUMN)
                        ),
                        (intro or None, new_founder or None, match.db_id),
                    )
                    current["founder"] = new_founder

                write_match_report(rows, matched, unmatched, fuzzy, backup_table)
                write_founder_report(founder_changes, founder_skipped)

                print(f"backup_table={backup_table}")
                print(f"excel_rows={len(rows)}")
                print(f"matched_records={len(matched)}")
                print(f"unmatched_rows={len(unmatched)}")
                print(f"founder_updates={len(founder_changes)}")
                print(f"match_report={MATCH_REPORT}")
                print(f"founder_report={FOUNDER_REPORT}")
    finally:
        conn.close()


if __name__ == "__main__":
    main()
