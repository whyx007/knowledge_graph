#!/usr/bin/env python3
"""Enrich PostgreSQL company fields from cleaned twopaper_v2.xlsx."""

from __future__ import annotations

import difflib
import re
from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

import psycopg2
from openpyxl import load_workbook
from psycopg2 import sql
from psycopg2.extras import RealDictCursor


ROOT = Path(__file__).resolve().parents[1]
XLSX_PATH = ROOT / "PostgreSQL " / "twopaper_v2.xlsx"
REPORT_PATH = ROOT / "docs" / "twopaper_v2_PostgreSQL字段丰富报告.md"
DSN = "postgresql://postgres:postgres@127.0.0.1:5432/ceo_brief"

FIELDS = ("core_tech", "products", "industry", "scenario", "cert_ip")


LEGAL_SUFFIXES = (
    "股份有限公司",
    "有限责任公司",
    "集团有限公司",
    "控股有限公司",
    "科技股份有限公司",
    "科技有限公司",
    "技术股份有限公司",
    "技术有限公司",
    "有限公司",
)

INDUSTRY_KEYWORDS = [
    "半导体",
    "集成电路",
    "先进封装",
    "功率器件",
    "射频元件",
    "光伏",
    "钙钛矿光伏",
    "新能源",
    "储能",
    "动力电池",
    "锂电",
    "汽车",
    "汽车电子",
    "船舶",
    "航天",
    "商业航天",
    "卫星通信",
    "卫星遥感",
    "卫星导航",
    "通信",
    "数据中心",
    "高性能计算",
    "机器人",
    "工业机器人",
    "人形机器人",
    "无人机",
    "低速无人车",
    "智能安防",
    "智能家居",
    "AR/VR",
    "医疗器械",
    "生物医药",
    "医药",
    "科研",
    "冶金",
    "化工",
    "石油石化",
    "核工业",
    "3C电子",
    "消费电子",
    "工业控制",
    "物流",
    "包装",
    "日化",
    "电力",
    "电子",
    "建筑光伏",
    "工商业分布式光伏",
]

SCENARIO_KEYWORDS = [
    "缺陷检测",
    "晶圆量测",
    "掩模版缺陷检测",
    "光罩制造",
    "晶圆衬底制造",
    "基板制造",
    "工业清洗",
    "表面处理",
    "精细微加工",
    "激光焊接",
    "激光切割",
    "发动机叶片打孔",
    "玻璃脆材加工",
    "SiC加工",
    "核工业去污",
    "低空目标发现识别跟踪与打击",
    "电路保护",
    "电流检测",
    "电源管理",
    "工业控制",
    "调压",
    "冷电联产",
    "节能减排",
    "空间深度感知",
    "人脸识别",
    "掌静脉识别",
    "物体识别",
    "空间建模",
    "避障导航",
    "定位导航",
    "物流拆垛",
    "机械臂引导",
    "目标追踪",
    "3D扫描建模",
    "星座布局与覆盖性分析",
    "卫星总体及分系统设计",
    "卫星全流程批量制造",
    "大型地面光伏电站",
    "建筑光伏一体化",
    "便携能源",
    "药物递送",
    "疫苗注射",
    "医美注射",
    "生物传感",
]

PRODUCT_WORD_RE = re.compile(
    r"(产品|设备|装备|仪器|系统|平台|模组|模块|器件|零部件|芯片|光源|激光器|相机|雷达|机器人|"
    r"解决方案|产线|材料|膜|元件|终端|服务|软件|整机|腔体|泵|机组)"
)
CORE_WORD_RE = re.compile(
    r"(技术|工艺|算法|模型|架构|芯片|材料|传感|激光|光学|控制|软件|平台|系统|制备|制造|检测|"
    r"测量|封装|设计|研发|开发|自主|专利|壁垒|能力|调控|合成|聚合|成像|通信|计算|训练|推理|编译)"
)
CERT_RE = re.compile(
    r"(专利|知识产权|自主知识产权|芯片设计IP|IP|认证|资质|高新技术企业|专精特新|AEC-Q100|安全认证|授权发明专利|国际发明专利)"
)
DROP_RE = re.compile(
    r"(融资|投资|估值|营收|收入|回款|订单|一供|独供|商业化进展|商业进展|股权|持股|上市|IPO|"
    r"获投|领投|跟投|资金支持|政府资源|市场验证|短板|风险|挑战|员工规模|人员规模|头部客户|批量导入|已导入|"
    r"股东|基金|地方国资|背靠|共同发起设立|资本|收益双方分成|参小股|出资建设|合资公司|客户资源)"
)
NOISE_PIECE_RE = re.compile(
    r"^(主打产品|核心产品|主要产品|产品及解决方案简介|产品简介|应用领域|应用场景|行业领域应用|产业链终端应用场景|"
    r"系统功能|公司荣誉资质|认证体系|专利|主打产品及应用领域|产品服务|技术特点)$"
)
FIELD_ROUTE_RE = re.compile(r"(目标市场领域|市场领域|应用领域|应用场景|用途|适用于|用于|面向)")
GENERIC_CLAIM_RE = re.compile(
    r"(新一轮技术革命|制高点|第一梯队|完美组合|人才密度|最前沿的认知|创领者|行业引领者|"
    r"国内外领先|原创技术|核心竞争力|无法替代|商业模式|公司团队|客户提供|市场潜力|爆发增长|"
    r"核心技术、平台生态、产品服务|平台型公司|平台公司)"
)
PROFILE_LINE_RE = re.compile(r"^(成立于|公司成立|是一家|公司定位|致力于|专注于|主要从事|公司从|公司通过|公司主要|公司拥有)")
GENERIC_PRODUCT_RE = re.compile(r"^(系统解决方案|全面的解决方案|完整解决方案|解决方案|配套管路系统|基于设备工程)")


@dataclass(frozen=True)
class Match:
    excel_name: str
    db_id: int
    db_name: str
    method: str


def clean_text(value: object) -> str:
    if value is None:
        return ""
    text = str(value).replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t\u3000]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def normalize_name(value: str) -> str:
    text = clean_text(value)
    text = text.replace("（", "(").replace("）", ")")
    text = re.sub(r"[-－—–](已退出|退出)$", "", text)
    text = re.sub(r"(已退出|退出)$", "", text)
    return re.sub(r"[\s\u3000]+", "", text)


def legal_core(value: str) -> str:
    text = normalize_name(value)
    for suffix in LEGAL_SUFFIXES:
        if text.endswith(suffix) and len(text) > len(suffix) + 2:
            return text[: -len(suffix)]
    return text


def split_company_names(value: str) -> list[str]:
    return [part.strip() for part in re.split(r"[、,，/]", clean_text(value)) if part.strip()]


def normalized_text(value: str) -> str:
    return re.sub(r"[\s，,。；;：:、（）()·.｜|]+", "", clean_text(value)).lower()


def strip_piece(value: str) -> str:
    text = clean_text(value)
    text = re.sub(r"^[•·\-—*]\s*", "", text)
    text = re.sub(r"^(?:[0-9]+|[一二三四五六七八九十]+)[.、）)]\s*", "", text)
    text = re.sub(
        r"^(核心技术|技术核心|核心产品|主要产品|公司产品|产品简介|产品及应用领域|应用领域|应用场景|"
        r"主打产品及应用领域|主打产品|产品服务|技术特点|系统功能|目标市场领域)[:：]\s*",
        "",
        text,
    )
    text = re.sub(r"^产品[0-9一二三四五六七八九十]+[:：]\s*", "", text)
    return text.strip(" 。；;")


def split_pieces(text: str) -> list[str]:
    text = clean_text(text)
    if not text:
        return []
    pieces: list[str] = []
    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        line = re.sub(r"(?<![A-Za-z0-9])([0-9]+[、)])", r"\n\1", line)
        line = re.sub(r"(核心技术[:：]|技术核心[:：]|核心产品[:：]|主要产品[:：]|产品简介[:：]|应用领域[:：]|应用场景[:：])", r"\n\1", line)
        pieces.extend(part.strip() for part in line.splitlines() if part.strip())
    return pieces


def split_sentences(text: str) -> list[str]:
    text = clean_text(text)
    if not text:
        return []
    parts = re.split(r"[。；;]\s*|\n+", text)
    return [strip_piece(part) for part in parts if strip_piece(part)]


def split_list_items(text: str) -> list[str]:
    text = clean_text(text)
    if not text:
        return []
    text = re.sub(r"[。；;]\s*", "\n", text)
    text = re.sub(r"(?<![A-Za-z0-9])(?:[0-9]+|[一二三四五六七八九十]+)[.、）)]", "\n", text)
    text = re.sub(r"[、,，]\s*(?=[\u4e00-\u9fffA-Za-z0-9]+(?:行业|领域|设备|系统|平台|材料|产品|电站|能源|检测|加工|制造|光伏|电池|机器人|数据中心|高性能计算))", "\n", text)
    return [strip_piece(part) for part in text.splitlines() if strip_piece(part)]


def drop_for_field(text: str, field: str) -> bool:
    text = strip_piece(text)
    if not text or NOISE_PIECE_RE.match(text) or DROP_RE.search(text):
        return True
    if re.match(r"^[\u4e00-\u9fff]。", text):
        return True
    if GENERIC_CLAIM_RE.search(text):
        return True
    if field == "core_tech":
        if FIELD_ROUTE_RE.search(text) or PROFILE_LINE_RE.search(text):
            return True
        if re.search(r"(销售|市场|客户|牵头承担|国家重点研发计划|关键技术重点专项)", text) and "技术" not in text:
            return True
        if not CORE_WORD_RE.search(text):
            return True
    elif field == "products":
        if PROFILE_LINE_RE.search(text) or text.startswith("公司"):
            return True
        if GENERIC_PRODUCT_RE.search(text):
            return True
        if re.search(r"(目标市场领域|应用领域|应用场景|用途[:：]?|适用于|用于|面向)", text) and not PRODUCT_WORD_RE.search(text):
            return True
        if not PRODUCT_WORD_RE.search(text):
            return True
    elif field == "cert_ip":
        if "暂无" in text or "未提及" in text or re.search(r"(服务的|客户为|客户包括|面向).*专精特新", text):
            return True
        if not CERT_RE.search(text):
            return True
    return False


def product_phrase(text: str) -> str:
    text = strip_piece(text)
    if "：" in text or ":" in text:
        head, tail = re.split(r"[:：]", text, 1)
        if PRODUCT_WORD_RE.search(head) and len(head) <= 50:
            return head.strip()
        if len(head) <= 18 and PRODUCT_WORD_RE.search(tail):
            text = tail.strip()
    if "提供" in text and PRODUCT_WORD_RE.search(text.split("提供", 1)[-1]):
        text = text.split("提供", 1)[-1]
    text = re.split(r"(?:，|,)?(?:用于|适用于|面向|应用于|服务于|可为|能够|实现|支持|满足|为)", text, 1)[0]
    text = re.sub(r"^(包括|主要包括|涵盖|产品包括)\s*", "", text)
    return strip_piece(text)


def core_phrase(text: str) -> str:
    text = strip_piece(text)
    if re.match(r"^[\u4e00-\u9fff]。", text):
        return ""
    if "天然气压差发电制冷系统" in text and re.search(r"研发、设计、生产、销售", text):
        return "天然气压差发电制冷系统研发与集成技术"
    if "透平膨胀机推动涡轮旋转带动发电机" in text:
        return "天然气透平膨胀发电制冷工艺"
    if "构建了" in text and "技术" in text:
        match = re.search(r"构建了(.+?技术)", text)
        if match:
            return strip_piece(match.group(1))
    if "深耕" in text and "技术" in text:
        match = re.search(r"深耕(.+?技术)", text)
        if match:
            return strip_piece(match.group(1))
    if re.search(r"研发、设计、生产、销售", text):
        return ""
    return text


def scenario_phrase(text: str) -> str:
    text = strip_piece(text)
    match = re.search(r"(?:用途[:：]|用于|适用于|应用于|面向|服务于)(.+)", text)
    if match:
        text = match.group(1)
    text = re.sub(r"^(用于|适用于|应用于|面向|各类|相关|主要|客户|核心客户)", "", text)
    return strip_piece(text)


def is_industry_only(text: str) -> bool:
    norm = normalized_text(text)
    if not norm:
        return True
    for keyword in INDUSTRY_KEYWORDS:
        key = normalized_text(keyword)
        if norm == key or norm in key or key in norm:
            return True
    return bool(re.search(r"(行业|领域|市场)$", text))


def append_unique(target: list[str], piece: str, max_items: int | None = None) -> bool:
    text = strip_piece(piece)
    if not text or DROP_RE.search(text) or NOISE_PIECE_RE.match(text):
        return False
    norm = normalized_text(text)
    if not norm:
        return False
    for idx, existing in enumerate(target):
        existing_norm = normalized_text(existing)
        if norm == existing_norm:
            return False
        if len(norm) >= 4 and norm in existing_norm:
            return False
        if len(existing_norm) >= 4 and existing_norm in norm:
            target[idx] = text
            return True
        if difflib.SequenceMatcher(None, norm, existing_norm).ratio() >= 0.9:
            if len(norm) > len(existing_norm):
                target[idx] = text
                return True
            return False
    if max_items is None or len(target) < max_items:
        target.append(text)
        return True
    return False


def unique_keywords(text: str, keywords: list[str]) -> list[str]:
    found: list[str] = []
    for keyword in keywords:
        if keyword in text and not any(keyword in old or old in keyword for old in found):
            found.append(keyword)
    return found


def extract_cert(text: str) -> list[str]:
    certs: list[str] = []
    for sentence in split_sentences(text):
        for item in cert_items(sentence):
            append_unique(certs, item, 8)
    return certs


def cert_items(sentence: str) -> list[str]:
    text = strip_piece(sentence)
    if not text or DROP_RE.search(text) or not CERT_RE.search(text):
        return []
    if "暂无" in text or "未提及" in text or re.search(r"(服务的|客户为|客户包括|面向).*专精特新", text):
        return []
    if re.search(r"(论文|期刊|会议|资助|基金|计划|项目|奖|人才|杰青|优青|千人|博物馆|新闻|进展)", text) and not re.search(
        r"(专利|知识产权|认证|资质|高新技术企业|专精特新)", text
    ):
        return []
    items: list[str] = []
    if "高新技术企业" in text:
        items.append("高新技术企业")
    if "专精特新" in text:
        match = re.search(r"[\u4e00-\u9fff]{0,8}专精特新[\u4e00-\u9fff]{0,4}", text)
        items.append(match.group(0) if match else "专精特新企业")
    for match in re.finditer(r"[^。；;，,]*?(?:AEC-Q100|BCTC|安全认证|认证|资质)[^。；;，,]*", text):
        value = match.group(0).strip(" ，,。；;")
        if value:
            items.append(value)
    for match in re.finditer(r"[^。；;，,]*?(?:自主知识产权|专利|授权发明专利|国际发明专利)[^。；;，,]*", text):
        value = match.group(0).strip(" ，,。；;")
        if value:
            items.append(value)
    return items or ([text] if len(text) <= 80 else [])


def extract_field_candidates(row: dict[str, str]) -> dict[str, list[str]]:
    intro = row.get("公司简介", "")
    core = row.get("核心技术", "")
    product = row.get("产品及应用领域", "")
    all_text = "\n".join([intro, core, product])

    out: dict[str, list[str]] = {field: [] for field in FIELDS}

    for piece in split_pieces(core):
        text = strip_piece(piece)
        if drop_for_field(text, "cert_ip") is False and CERT_RE.search(text):
            for item in cert_items(text):
                append_unique(out["cert_ip"], item, 8)
        else:
            core_text = core_phrase(text)
            if not drop_for_field(core_text, "core_tech"):
                append_unique(out["core_tech"], core_text, 8)

    for piece in split_pieces(product):
        text = strip_piece(piece)
        if drop_for_field(text, "cert_ip") is False and CERT_RE.search(text):
            for item in cert_items(text):
                append_unique(out["cert_ip"], item, 8)
            continue
        core_text = core_phrase(text)
        if re.match(r"^(技术核心|核心技术)[:：]", piece.strip()) and not drop_for_field(core_text, "core_tech"):
            append_unique(out["core_tech"], core_text, 8)
            continue
        if PRODUCT_WORD_RE.search(text):
            product_text = product_phrase(text)
            if not drop_for_field(product_text, "products"):
                append_unique(out["products"], product_text, 10)
        if FIELD_ROUTE_RE.search(text):
            scenario_text = scenario_phrase(text)
            for item in split_list_items(scenario_text):
                if (
                    2 <= len(item) <= 50
                    and not DROP_RE.search(item)
                    and not PRODUCT_WORD_RE.search(item)
                    and not is_industry_only(item)
                ):
                    append_unique(out["scenario"], item, 12)

    for sentence in split_sentences(intro):
        if drop_for_field(sentence, "cert_ip") is False and CERT_RE.search(sentence):
            for item in cert_items(sentence):
                append_unique(out["cert_ip"], item, 8)
        elif re.search(r"(掌握|自主研发|核心技术|深耕.*技术|构建.*技术|突破.*技术|技术栈|技术体系)", sentence):
            core_text = core_phrase(sentence)
            if not drop_for_field(core_text, "core_tech"):
                append_unique(out["core_tech"], core_text, 8)
        elif re.search(r"(主要产品|核心产品|主打产品|产品包括|产品涵盖|提供.*(?:产品|解决方案|服务)|解决方案包括)", sentence):
            product_text = product_phrase(sentence)
            if not drop_for_field(product_text, "products"):
                append_unique(out["products"], product_text, 10)

    for keyword in unique_keywords(all_text, INDUSTRY_KEYWORDS):
        append_unique(out["industry"], keyword, 12)
    for keyword in unique_keywords(all_text, SCENARIO_KEYWORDS):
        append_unique(out["scenario"], keyword, 12)

    return out


def split_existing(value: str, field: str | None = None) -> list[str]:
    text = clean_text(value)
    if not text:
        return []
    parts: list[str] = []
    for part in re.split(r"\n+|；|;(?=\s*[\u4e00-\u9fffA-Za-z0-9])", text):
        stripped = strip_piece(part)
        if stripped:
            if NOISE_PIECE_RE.match(stripped):
                continue
            if field == "cert_ip":
                for sentence in split_sentences(stripped):
                    for item in cert_items(sentence):
                        if item and not drop_for_field(item, "cert_ip"):
                            parts.append(item)
                continue
            if field in {"core_tech", "products"} and drop_for_field(stripped, field):
                continue
            if field in {"industry", "scenario"} and (DROP_RE.search(stripped) or GENERIC_CLAIM_RE.search(stripped)):
                continue
            parts.append(stripped)
    if field == "cert_ip":
        return parts
    return parts or [text]


def merge_field(existing: str, additions: list[str], field: str, max_total: int = 18) -> tuple[str, list[str]]:
    pieces = split_existing(existing, field)
    before_norm = normalized_text("；".join(pieces))
    added: list[str] = []
    for addition in additions:
        old = list(pieces)
        if append_unique(pieces, addition, max_total):
            if normalized_text("；".join(pieces)) != normalized_text("；".join(old)):
                added.append(strip_piece(addition))
    if not pieces:
        merged = ""
    elif len(pieces) == 1:
        merged = pieces[0]
    else:
        merged = "；".join(pieces)
    return merged, added if normalized_text(merged) != before_norm else []


def load_excel_rows() -> list[dict[str, str]]:
    wb = load_workbook(XLSX_PATH, read_only=True, data_only=True)
    ws = wb["企业信息提取"]
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    rows = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        row = {key: clean_text(value) for key, value in zip(headers, values)}
        if row.get("公司名称"):
            rows.append(row)
    return rows


def unique_core_index(companies: list[dict]) -> dict[str, dict]:
    buckets: dict[str, list[dict]] = defaultdict(list)
    for company in companies:
        buckets[legal_core(company["name"])].append(company)
    return {key: values[0] for key, values in buckets.items() if len(values) == 1}


def find_matches(rows: list[dict[str, str]], companies: list[dict]) -> tuple[list[tuple[dict, Match]], list[dict]]:
    by_name = {company["name"]: company for company in companies}
    by_norm = {normalize_name(company["name"]): company for company in companies}
    by_core = unique_core_index(companies)
    matched: list[tuple[dict, Match]] = []
    unmatched: list[dict] = []
    seen_pairs: set[tuple[str, int]] = set()

    for row in rows:
        excel_name = row["公司名称"]
        candidates: list[Match] = []
        if excel_name in by_name:
            c = by_name[excel_name]
            candidates.append(Match(excel_name, c["id"], c["name"], "exact"))
        elif normalize_name(excel_name) in by_norm:
            c = by_norm[normalize_name(excel_name)]
            candidates.append(Match(excel_name, c["id"], c["name"], "normalized_parentheses_space"))
        else:
            for part in split_company_names(excel_name):
                if normalize_name(part) in by_norm:
                    c = by_norm[normalize_name(part)]
                    candidates.append(Match(part, c["id"], c["name"], "split_exact"))
            if not candidates and legal_core(excel_name) in by_core:
                c = by_core[legal_core(excel_name)]
                candidates.append(Match(excel_name, c["id"], c["name"], "legal_suffix"))

        if candidates:
            for candidate in candidates:
                pair = (excel_name, candidate.db_id)
                if pair not in seen_pairs:
                    matched.append((row, candidate))
                    seen_pairs.add(pair)
        else:
            unmatched.append(row)
    return matched, unmatched


def md_cell(value: object, limit: int = 180) -> str:
    text = re.sub(r"\s+", " ", "" if value is None else str(value)).strip()
    if len(text) > limit:
        text = text[: limit - 1] + "…"
    return text.replace("|", "\\|")


def write_report(backup_table: str, matched_count: int, unmatched_count: int, changes: list[dict]) -> None:
    field_counts = {field: sum(1 for change in changes if change["added"].get(field)) for field in FIELDS}
    lines = [
        "# twopaper_v2.xlsx 丰富 PostgreSQL companies 字段报告",
        "",
        f"- 生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
        f"- Excel 文件：`{XLSX_PATH}`",
        "- 数据库：`ceo_brief.public.companies`",
        f"- 更新前备份表：`{backup_table}`",
        f"- 高置信匹配记录：{matched_count}",
        f"- 未写库 Excel 行：{unmatched_count}",
        f"- 实际更新企业：{len(changes)}",
        "",
        "## 字段更新统计",
        "",
        "| 字段 | 更新企业数 |",
        "|---|---:|",
    ]
    for field in FIELDS:
        lines.append(f"| {field} | {field_counts[field]} |")

    lines.extend(
        [
            "",
            "## 字段边界口径",
            "",
            "- `core_tech`：只补充技术、工艺、算法、架构、研发/制造/检测能力等。",
            "- `products`：只补充产品、设备、模组、器件、平台、软件、解决方案等。",
            "- `industry`：只补充行业/领域名，例如半导体、光伏、航天、数据中心等。",
            "- `scenario`：只补充具体使用场景/工艺场景，例如缺陷检测、激光清洗、物流拆垛等。",
            "- `cert_ip`：只补充专利、知识产权、认证、资质等明确表述。",
            "",
            "## 更新明细",
            "",
            "| DB id | 公司名称 | 匹配方式 | 更新字段 | 新增内容摘要 |",
            "|---:|---|---|---|---|",
        ]
    )
    for change in changes:
        for field in FIELDS:
            added = change["added"].get(field) or []
            if added:
                lines.append(
                    f"| {change['db_id']} | {md_cell(change['db_name'], 80)} | {change['method']} | "
                    f"{field} | {md_cell('；'.join(added), 260)} |"
                )
    REPORT_PATH.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> None:
    rows = load_excel_rows()
    backup_table = f"companies_twopaper_field_enrich_backup_{datetime.now().strftime('%Y%m%d_%H%M%S')}"

    conn = psycopg2.connect(DSN)
    try:
        with conn:
            with conn.cursor(cursor_factory=RealDictCursor) as cur:
                cur.execute("SELECT id, name, core_tech, products, industry, scenario, cert_ip FROM companies ORDER BY id")
                companies = list(cur.fetchall())
                matched, unmatched = find_matches(rows, companies)

                cur.execute(
                    sql.SQL(
                        """
                        CREATE TABLE {} AS
                        SELECT id, name, core_tech, products, industry, scenario, cert_ip
                        FROM companies
                        """
                    ).format(sql.Identifier(backup_table))
                )

                by_id = {company["id"]: company for company in companies}
                changes: list[dict] = []

                for row, match in matched:
                    current = by_id[match.db_id]
                    candidates = extract_field_candidates(row)
                    new_values: dict[str, str] = {}
                    added_by_field: dict[str, list[str]] = {}

                    for field in FIELDS:
                        max_total = 24 if field in {"products", "scenario"} else 18
                        merged, added = merge_field(clean_text(current.get(field, "")), candidates[field], field, max_total=max_total)
                        new_values[field] = merged or None
                        if added:
                            added_by_field[field] = added

                    if added_by_field:
                        cur.execute(
                            """
                            UPDATE companies
                            SET core_tech = %(core_tech)s,
                                products = %(products)s,
                                industry = %(industry)s,
                                scenario = %(scenario)s,
                                cert_ip = %(cert_ip)s
                            WHERE id = %(id)s
                            """,
                            {
                                "id": match.db_id,
                                "core_tech": new_values["core_tech"],
                                "products": new_values["products"],
                                "industry": new_values["industry"],
                                "scenario": new_values["scenario"],
                                "cert_ip": new_values["cert_ip"],
                            },
                        )
                        changes.append(
                            {
                                "db_id": match.db_id,
                                "db_name": match.db_name,
                                "method": match.method,
                                "added": added_by_field,
                            }
                        )

                write_report(backup_table, len(matched), len(unmatched), changes)
                print(f"backup_table={backup_table}")
                print(f"matched={len(matched)}")
                print(f"unmatched={len(unmatched)}")
                print(f"updated_companies={len(changes)}")
                for field in FIELDS:
                    print(f"{field}_updates={sum(1 for change in changes if change['added'].get(field))}")
                print(f"report={REPORT_PATH}")
    finally:
        conn.close()


if __name__ == "__main__":
    main()
