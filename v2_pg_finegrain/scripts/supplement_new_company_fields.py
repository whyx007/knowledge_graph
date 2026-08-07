#!/usr/bin/env python3
"""Fill missing fields for the 46 newly added companies without overwriting data."""

from __future__ import annotations

import argparse
import json
import os
import shutil
import subprocess
import sys
from collections import Counter
from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path
from typing import Any
from uuid import uuid4

import psycopg2
from neo4j import GraphDatabase
from openpyxl import load_workbook
from psycopg2.extras import RealDictCursor

from apply_manual_company_updates import (
    DEFAULT_BACKUP_ROOT,
    DEFAULT_DSN,
    neo4j_snapshot,
    pg_snapshot,
    restore_neo4j,
    sha256_file,
    write_json,
)


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_XLSX = ROOT / "PostgreSQL " / "遗漏被投企业名单.xlsx"
EXPECTED_COMPANIES = 526
EXPECTED_SOURCE_ROWS = 46
EXPECTED_MOUNTS = 293
NEW_ID_MIN = 1441
NEW_ID_MAX = 1486

TEXT_FIELDS = ("domain", "industry", "core_tech", "cert_ip", "status", "investor", "city")
NUMERIC_FIELDS = ("valuation", "investment")
ALL_FIELDS = (*TEXT_FIELDS, *NUMERIC_FIELDS)

SOURCE_MAP = {
    "行业标签": "domain",
    "应用标签": "industry",
    "技术标签": "core_tech",
    "企业认定标签": "cert_ip",
    "持有状态": "status",
    "投资项目负责人（最新一轮）": "investor",
    "注册地址": "city",
    "最新估值（万元）": "valuation",
    "投资总额（万元）": "investment",
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="只补充新增企业的空字段")
    parser.add_argument("--xlsx", type=Path, default=DEFAULT_XLSX)
    parser.add_argument("--dsn", default=os.getenv("PG_DSN", DEFAULT_DSN))
    parser.add_argument("--neo4j-uri", default=os.getenv("NEO4J_URI", "bolt://127.0.0.1:7688"))
    parser.add_argument("--neo4j-user", default=os.getenv("NEO4J_USER", "neo4j"))
    parser.add_argument("--neo4j-password", default=os.getenv("NEO4J_PASSWORD", "neo4j2026"))
    parser.add_argument("--backup-root", type=Path, default=DEFAULT_BACKUP_ROOT)
    parser.add_argument("--apply", action="store_true", help="备份后执行双库写入")
    return parser.parse_args()


def clean(value: Any) -> str | None:
    if value is None:
        return None
    text = " ".join(str(value).replace("\r", " ").replace("\n", " ").split()).strip()
    return text or None


def normalize_city(value: Any) -> str | None:
    text = clean(value)
    if not text:
        return None
    parts = [part.strip() for part in text.split("-") if part.strip()]
    if len(parts) >= 2 and parts[0] == parts[1]:
        parts.pop(0)
    return ",".join(parts)


def normalize_investor(value: Any) -> str | None:
    text = clean(value)
    if not text:
        return None
    names = [name.strip() for name in text.replace("，", ",").split(",") if name.strip()]
    return " ".join(f"{index}.{name}" for index, name in enumerate(names, start=1))


def normalize_numeric(value: Any) -> Decimal | None:
    if value is None or str(value).strip() == "":
        return None
    return Decimal(str(value).replace(",", "").strip())


def load_source(path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    if "遗漏被投企业名单" not in workbook.sheetnames:
        raise RuntimeError("缺少工作表：遗漏被投企业名单")
    sheet = workbook["遗漏被投企业名单"]
    headers = [clean(value) for value in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))]
    required = {"客户名称", *SOURCE_MAP.keys()}
    missing = sorted(required - set(headers))
    if missing:
        raise RuntimeError(f"源表缺少字段：{missing}")
    indexes = {header: index for index, header in enumerate(headers)}
    rows: list[dict[str, Any]] = []
    for excel_row, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        name = clean(values[indexes["客户名称"]])
        if not name:
            continue
        row: dict[str, Any] = {"excel_row": excel_row, "name": name}
        for source_column, db_column in SOURCE_MAP.items():
            value = values[indexes[source_column]]
            if db_column == "city":
                value = normalize_city(value)
            elif db_column == "investor":
                value = normalize_investor(value)
            elif db_column in NUMERIC_FIELDS:
                value = normalize_numeric(value)
            else:
                value = clean(value)
            row[db_column] = value
        rows.append(row)
    counts = Counter(row["name"] for row in rows)
    duplicates = sorted(name for name, count in counts.items() if count > 1)
    if duplicates:
        raise RuntimeError(f"源表企业名称重复：{duplicates}")
    if len(rows) != EXPECTED_SOURCE_ROWS:
        raise RuntimeError(f"源表企业数应为 {EXPECTED_SOURCE_ROWS}，实际 {len(rows)}")
    return rows


def fetch_targets(dsn: str, names: list[str]) -> list[dict[str, Any]]:
    connection = psycopg2.connect(dsn)
    try:
        connection.autocommit = False
        with connection.cursor(cursor_factory=RealDictCursor) as cursor:
            cursor.execute("SET TRANSACTION READ ONLY")
            cursor.execute(
                "SELECT id,name,domain,industry,core_tech,cert_ip,status,investor,city,"
                "valuation,investment FROM public.companies WHERE name = ANY(%s) ORDER BY id",
                (names,),
            )
            rows = [dict(row) for row in cursor.fetchall()]
            cursor.execute("SELECT count(*) AS count FROM public.companies")
            total = cursor.fetchone()["count"]
        connection.rollback()
    finally:
        connection.close()
    if total != EXPECTED_COMPANIES:
        raise RuntimeError(f"PostgreSQL 企业数应为 {EXPECTED_COMPANIES}，实际 {total}")
    if len(rows) != EXPECTED_SOURCE_ROWS:
        found = {row["name"] for row in rows}
        raise RuntimeError(f"PostgreSQL 未精确匹配：{sorted(set(names) - found)}")
    ids = {row["id"] for row in rows}
    expected_ids = set(range(NEW_ID_MIN, NEW_ID_MAX + 1))
    if ids != expected_ids:
        raise RuntimeError("匹配企业不是上一批新增的 PostgreSQL id 1441-1486")
    return rows


def is_missing(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def build_plan(
    source_rows: list[dict[str, Any]], target_rows: list[dict[str, Any]]
) -> tuple[list[dict[str, Any]], dict[str, int]]:
    source_by_name = {row["name"]: row for row in source_rows}
    plan: list[dict[str, Any]] = []
    additions = Counter()
    for target in target_rows:
        source = source_by_name[target["name"]]
        final = {"id": target["id"], "name": target["name"]}
        changed_fields: list[str] = []
        for field in ALL_FIELDS:
            old = target[field]
            candidate = source[field]
            if is_missing(old) and candidate is not None:
                final[field] = candidate
                changed_fields.append(field)
                additions[field] += 1
            else:
                final[field] = old
        final["changed_fields"] = changed_fields
        plan.append(final)
    return plan, dict(sorted(additions.items()))


def create_backups(
    args: argparse.Namespace,
    source_path: Path,
    source_sha256: str,
    neo_before: dict[str, Any],
    plan: list[dict[str, Any]],
    additions: dict[str, int],
    run_id: str,
) -> Path:
    backup_dir = args.backup_root.resolve() / run_id
    backup_dir.mkdir(parents=True, exist_ok=False)
    shutil.copy2(source_path, backup_dir / "source.xlsx")
    dump_path = backup_dir / "postgres_companies_before.dump"
    subprocess.run(
        [
            "pg_dump", "--format=custom", "--no-owner", "--no-privileges",
            "--table=public.companies", f"--file={dump_path}", args.dsn,
        ],
        check=True,
        capture_output=True,
        text=True,
    )
    connection = psycopg2.connect(args.dsn)
    try:
        connection.autocommit = False
        with connection.cursor() as cursor:
            cursor.execute("SET TRANSACTION READ ONLY")
        postgres_before = pg_snapshot(connection)
        connection.rollback()
    finally:
        connection.close()
    write_json(backup_dir / "postgres_companies_before.json", postgres_before)
    write_json(backup_dir / "neo4j_enterprises_before.json", neo_before["enterprises"])
    write_json(
        backup_dir / "neo4j_enterprise_relationships_before.json",
        neo_before["enterprise_relationships"],
    )
    write_json(
        backup_dir / "preflight.json",
        {
            "run_id": run_id,
            "source_sha256": source_sha256,
            "source_rows": len(plan),
            "additions_by_field": additions,
        },
    )
    manifest = {
        path.name: {"size": path.stat().st_size, "sha256": sha256_file(path)}
        for path in backup_dir.iterdir()
        if path.is_file()
    }
    write_json(backup_dir / "manifest.json", manifest)
    return backup_dir


def apply_postgres(connection: Any, plan: list[dict[str, Any]]) -> list[dict[str, Any]]:
    with connection.cursor(cursor_factory=RealDictCursor) as cursor:
        cursor.execute("LOCK TABLE public.companies IN SHARE ROW EXCLUSIVE MODE")
        cursor.execute("SELECT count(*) AS count FROM public.companies")
        if cursor.fetchone()["count"] != EXPECTED_COMPANIES:
            raise RuntimeError("PostgreSQL 企业数在预检后发生变化")
        assignments = ", ".join(
            f"{field} = CASE WHEN {field} IS NULL"
            + (f" OR btrim({field}) = ''" if field in TEXT_FIELDS else "")
            + f" THEN %s ELSE {field} END"
            for field in ALL_FIELDS
        )
        sql = f"UPDATE public.companies SET {assignments} WHERE id = %s"
        for row in plan:
            cursor.execute(sql, [*(row[field] for field in ALL_FIELDS), row["id"]])
            if cursor.rowcount != 1:
                raise RuntimeError(f"PostgreSQL id={row['id']} 补充失败")
        cursor.execute(
            "SELECT id,name,domain,industry,core_tech,cert_ip,status,investor,city,"
            "valuation,investment FROM public.companies WHERE id BETWEEN %s AND %s ORDER BY id",
            (NEW_ID_MIN, NEW_ID_MAX),
        )
        return [dict(row) for row in cursor.fetchall()]


def neo_payload(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    payload = []
    for row in rows:
        item = {"enterprise_id": f"ENT_PG_{row['id']}"}
        for field in ALL_FIELDS:
            value = row[field]
            if isinstance(value, Decimal):
                value = float(value)
            item[field] = value
        payload.append(item)
    return payload


def apply_neo4j(driver: Any, rows: list[dict[str, Any]], run_id: str) -> None:
    text_assignments = ", ".join(
        f"e.{field} = CASE WHEN e.{field} IS NULL OR trim(e.{field}) = '' "
        f"THEN row.{field} ELSE e.{field} END"
        for field in TEXT_FIELDS
    )
    numeric_assignments = ", ".join(
        f"e.{field} = CASE WHEN e.{field} IS NULL THEN row.{field} ELSE e.{field} END"
        for field in NUMERIC_FIELDS
    )
    query = (
        "UNWIND $rows AS row MATCH (e:Enterprise {enterprise_id: row.enterprise_id}) SET "
        + text_assignments
        + ", "
        + numeric_assignments
        + ", e.supplement_source = 'missing_invested_companies_xlsx', "
        "e.supplement_run_id = $run_id, e.last_supplement_at = datetime()"
    )
    with driver.session() as session:
        transaction = session.begin_transaction()
        try:
            transaction.run(query, rows=rows, run_id=run_id).consume()
            counts = transaction.run(
                "MATCH (e:Enterprise) WITH count(e) AS enterprises "
                "MATCH ()-[r:LOCATED_IN_SUBSTAGE]->() "
                "RETURN enterprises, count(r) AS mounts"
            ).single(strict=True)
            if counts["enterprises"] != EXPECTED_COMPANIES or counts["mounts"] != EXPECTED_MOUNTS:
                raise RuntimeError(f"Neo4j 事务内计数异常：{dict(counts)}")
            supplemented = transaction.run(
                "MATCH (e:Enterprise {supplement_run_id: $run_id}) RETURN count(e) AS count",
                run_id=run_id,
            ).single(strict=True)["count"]
            if supplemented != EXPECTED_SOURCE_ROWS:
                raise RuntimeError(f"Neo4j 仅补充 {supplemented} 个节点")
            transaction.commit()
        except Exception:
            transaction.rollback()
            raise


def verify(
    args: argparse.Namespace,
    before: list[dict[str, Any]],
    expected: list[dict[str, Any]],
    relationships_before: list[dict[str, Any]],
    run_id: str,
) -> dict[str, Any]:
    actual = fetch_targets(args.dsn, [row["name"] for row in expected])
    expected_by_id = {row["id"]: row for row in expected}
    before_by_id = {row["id"]: row for row in before}
    pg_mismatches = []
    overwritten = []
    for row in actual:
        wanted = expected_by_id[row["id"]]
        old = before_by_id[row["id"]]
        for field in ALL_FIELDS:
            if row[field] != wanted[field]:
                pg_mismatches.append((row["id"], field))
            if not is_missing(old[field]) and row[field] != old[field]:
                overwritten.append((row["id"], field))

    driver = GraphDatabase.driver(
        args.neo4j_uri, auth=(args.neo4j_user, args.neo4j_password)
    )
    try:
        with driver.session() as session:
            neo_rows = [
                dict(row)
                for row in session.run(
                    "MATCH (e:Enterprise) WHERE toInteger(e.source_pk) >= $low "
                    "AND toInteger(e.source_pk) <= $high "
                    "RETURN e.enterprise_id AS enterprise_id, e.domain AS domain, "
                    "e.industry AS industry, e.core_tech AS core_tech, e.cert_ip AS cert_ip, "
                    "e.status AS status, e.investor AS investor, e.city AS city, "
                    "e.valuation AS valuation, e.investment AS investment ORDER BY enterprise_id",
                    low=NEW_ID_MIN,
                    high=NEW_ID_MAX,
                )
            ]
            relationships_after = [
                {
                    "enterprise_id": row["enterprise_id"],
                    "type": row["type"],
                    "other_labels": row["other_labels"],
                    "other_identity": row["other_identity"],
                    "properties": dict(row["properties"]),
                }
                for row in session.run(
                    "MATCH (e:Enterprise)-[r]-(other) "
                    "RETURN e.enterprise_id AS enterprise_id, type(r) AS type, "
                    "labels(other) AS other_labels, "
                    "coalesce(other.substage_id,other.enterprise_id,other.stage_id,"
                    "other.segment_id,other.chain_id) AS other_identity, properties(r) AS properties "
                    "ORDER BY enterprise_id,type,other_identity"
                )
            ]
            run_count = session.run(
                "MATCH (e:Enterprise {supplement_run_id: $run_id}) RETURN count(e) AS count",
                run_id=run_id,
            ).single(strict=True)["count"]
    finally:
        driver.close()

    neo_mismatches = []
    for row in neo_rows:
        company_id = int(row["enterprise_id"].removeprefix("ENT_PG_"))
        wanted = expected_by_id[company_id]
        for field in ALL_FIELDS:
            expected_value = wanted[field]
            actual_value = row[field]
            if field in NUMERIC_FIELDS and expected_value is not None:
                if Decimal(str(actual_value)) != expected_value:
                    neo_mismatches.append((company_id, field))
            elif actual_value != expected_value:
                neo_mismatches.append((company_id, field))
    if pg_mismatches or overwritten or neo_mismatches:
        raise RuntimeError(
            f"验收失败：pg={pg_mismatches[:10]}, overwritten={overwritten[:10]}, "
            f"neo4j={neo_mismatches[:10]}"
        )
    if relationships_after != relationships_before:
        raise RuntimeError("企业关系或关系属性发生变化")
    if len(neo_rows) != EXPECTED_SOURCE_ROWS or run_count != EXPECTED_SOURCE_ROWS:
        raise RuntimeError("Neo4j 补充节点数验收失败")
    return {
        "postgres_companies_checked": len(actual),
        "neo4j_enterprises_checked": len(neo_rows),
        "existing_values_overwritten": 0,
        "postgres_field_mismatches": 0,
        "neo4j_field_mismatches": 0,
        "enterprise_relationships_unchanged": len(relationships_after),
    }


def main() -> None:
    args = parse_args()
    source_path = args.xlsx.resolve()
    source_rows = load_source(source_path)
    before = fetch_targets(args.dsn, [row["name"] for row in source_rows])
    plan, additions = build_plan(source_rows, before)

    driver = GraphDatabase.driver(
        args.neo4j_uri, auth=(args.neo4j_user, args.neo4j_password)
    )
    try:
        neo_before = neo4j_snapshot(driver)
    finally:
        driver.close()
    if neo_before["counts"] != {"enterprises": EXPECTED_COMPANIES, "mounts": EXPECTED_MOUNTS}:
        raise RuntimeError(f"Neo4j 基线异常：{neo_before['counts']}")

    source_sha256 = sha256_file(source_path)
    preview = {
        "source_sha256": source_sha256,
        "matched_new_companies": len(plan),
        "additions_by_field": additions,
        "ignored_source_columns": ["是否追投", "客户分类", "首投时间", "科服项目", "备注"],
        "overwrite_policy": "only_fill_null_or_empty",
    }
    print(json.dumps(preview, ensure_ascii=False, indent=2))
    if not args.apply:
        print("\n预检通过；未执行写入。使用 --apply 执行。")
        return

    run_id = "supplement_new_companies_" + datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ") + "_" + uuid4().hex[:8]
    backup_dir = create_backups(
        args, source_path, source_sha256, neo_before, plan, additions, run_id
    )
    connection = psycopg2.connect(args.dsn)
    driver = GraphDatabase.driver(
        args.neo4j_uri, auth=(args.neo4j_user, args.neo4j_password)
    )
    neo_committed = False
    try:
        connection.autocommit = False
        final_rows = apply_postgres(connection, plan)
        apply_neo4j(driver, neo_payload(final_rows), run_id)
        neo_committed = True
        connection.commit()
    except Exception:
        connection.rollback()
        if neo_committed:
            restore_neo4j(driver, neo_before["enterprises"], [])
        raise
    finally:
        connection.close()
        driver.close()

    verification = verify(
        args,
        before,
        final_rows,
        neo_before["enterprise_relationships"],
        run_id,
    )
    result = {
        **preview,
        "run_id": run_id,
        "backup_dir": str(backup_dir),
        "verification": verification,
        "completed_at_utc": datetime.now(timezone.utc).isoformat(timespec="seconds"),
    }
    write_json(backup_dir / "result.json", result)
    print("\n" + json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        raise
